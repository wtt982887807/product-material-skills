"""Read-only business-document inventory with explicit evidence and parse boundaries.

No document contents are executed. Source files are never saved or converted here.
This module produces candidates, not certified product identities or legal opinions.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
import logging
import os
from pathlib import Path
import re
import sys
import unicodedata
from zipfile import ZipFile

from lxml import etree

SUPPORTED = {'.docx', '.pptx', '.pdf', '.xlsx', '.doc', '.xls', '.png', '.jpg', '.jpeg'}
NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}
EXCLUDED_DIRS = {
    'node_modules', '__pycache__', 'venv', 'dist', 'build', 'qa', 'tmp', 'output', 'outputs',
    '交付文件', '验证材料', '最终检查', '版式检查', '留用答辩ppt_制作过程',
}


def _xml(raw):
    return etree.fromstring(raw, parser=etree.XMLParser(resolve_entities=False, no_network=True))


def _tag(prefix, name):
    return '{' + NS[prefix] + '}' + name


def _has_deleted_ancestor(node):
    return bool(node.xpath('ancestor-or-self::w:del | ancestor-or-self::w:moveFrom', namespaces=NS))


def _word_text(node):
    chunks = []
    for child in node.iter():
        if _has_deleted_ancestor(child):
            continue
        if child.tag == _tag('w', 't'):
            chunks.append(child.text or '')
        elif child.tag == _tag('w', 'tab'):
            chunks.append('\t')
        elif child.tag in {_tag('w', 'br'), _tag('w', 'cr')}:
            chunks.append('\n')
    return ''.join(chunks).strip()


def _heading_level(p, text):
    if not text or len(text) > 120:
        return None
    outline = p.find('w:pPr/w:outlineLvl', NS)
    if outline is not None:
        return min(9, int(outline.get(_tag('w', 'val'), 0)) + 1)
    style = p.find('w:pPr/w:pStyle', NS)
    if style is not None:
        match = re.search(r'(?:heading|标题)\s*([1-9])', style.get(_tag('w', 'val'), ''), re.I)
        if match:
            return int(match.group(1))
    if re.match(r'^第[〇零一二三四五六七八九十百0-9]+(?:部分|章)', text):
        return 1
    if re.match(r'^[一二三四五六七八九十]+[、．.]', text):
        return 2
    if re.match(r'^[（(][一二三四五六七八九十]+[）)]', text):
        return 3
    return None


def _result(blocks, status='parsed', issues=None, **extra):
    return {'blocks': blocks, 'parse_status': status, 'issues': issues or [],
            'revision_state': 'not_detected', **extra}


def parse_docx(path, doc_id):
    blocks, issues = [], []
    revisions = False
    with ZipFile(path) as z:
        names = set(z.namelist())
        parts = ['word/document.xml']
        parts += sorted(n for n in names if re.fullmatch(r'word/(?:header|footer)\d+\.xml', n))
        parts += [n for n in ('word/footnotes.xml', 'word/endnotes.xml') if n in names]
        if 'word/comments.xml' in names:
            issues.append('comments_present_excluded_from_current_body')
        for part in parts:
            if part not in names:
                continue
            root = _xml(z.read(part))
            if root.xpath('//w:ins | //w:del | //w:moveFrom | //w:moveTo | //w:pPrChange | //w:rPrChange | //w:tblPrChange | //w:cellDel | //w:cellIns | //w:cellMerge', namespaces=NS):
                revisions = True
            paragraphs = {p: i for i, p in enumerate(root.xpath('.//w:p', namespaces=NS), 1)}
            tables = {t: i for i, t in enumerate(root.xpath('.//w:tbl', namespaces=NS), 1)}
            for alternate in root.xpath('//*[local-name()="AlternateContent"]'):
                choices = [c for c in alternate if etree.QName(c).localname == 'Choice']
                fallbacks = [c for c in alternate if etree.QName(c).localname == 'Fallback']
                selected = choices[0] if choices else (fallbacks[0] if fallbacks else None)
                for child in list(alternate):
                    if child is not selected:
                        alternate.remove(child)
                issues.append('compatibility_alternative_first_choice_selected_not_double_counted')
            section = []

            def emit(kind, text, location, **extra):
                if text.strip():
                    anchor = 'p' + str(location['paragraph']) if 'paragraph' in location else 't' + str(location['table']) + 'r' + str(location['row'])
                    blocks.append({'doc_id': doc_id, 'block_id': f'{doc_id}:{part}:{anchor}',
                                   'kind': kind, 'text': text, 'location': {'part': part, **location},
                                   'section_path': list(section), **extra})

            def walk(node):
                nonlocal section
                if _has_deleted_ancestor(node):
                    return
                if node.tag == _tag('w', 'p'):
                    # Paragraphs within text boxes remain anchored within their source paragraph.
                    text = _word_text(node)
                    level = _heading_level(node, text)
                    if level:
                        section = section[:level - 1] + [text]
                    emit('heading' if level else 'paragraph', text, {'paragraph': paragraphs[node]},
                         revision_state='inserted_text_present' if node.xpath('.//w:ins | .//w:moveTo', namespaces=NS) else 'current_view')
                    return
                if node.tag == _tag('w', 'tbl'):
                    table_index = tables[node]
                    rows = [r for r in node.xpath('.//w:tr', namespaces=NS)
                            if r.xpath('ancestor::w:tbl[1]', namespaces=NS)[0] is node]
                    for row_index, row in enumerate(rows, 1):
                        if _has_deleted_ancestor(row) or row.find('w:trPr/w:del', NS) is not None:
                            continue
                        cells, locations = [], []
                        for col, cell in enumerate(row.findall('w:tc', NS), 1):
                            if cell.find('w:tcPr/w:cellDel', NS) is not None:
                                continue
                            ps = [p for p in cell.xpath('.//w:p', namespaces=NS)
                                  if p.xpath('ancestor::w:tbl[1]', namespaces=NS)[0] is node]
                            texts = [_word_text(p) for p in ps]
                            cells.append('\n'.join(t for t in texts if t))
                            loc = {'part': part, 'table': table_index, 'row': row_index, 'column': col,
                                   'paragraphs': [paragraphs[p] for p in ps]}
                            grid = cell.find('w:tcPr/w:gridSpan', NS)
                            merge = cell.find('w:tcPr/w:vMerge', NS)
                            if grid is not None:
                                loc['grid_span'] = grid.get(_tag('w', 'val'))
                            if merge is not None:
                                loc['vertical_merge'] = merge.get(_tag('w', 'val'), 'continue')
                            locations.append(loc)
                        emit('table_row', ' | '.join(cells), {'table': table_index, 'row': row_index},
                             cells=cells, cell_locations=locations)
                    for nested in node.xpath('.//w:tbl', namespaces=NS):
                        if nested.xpath('ancestor::w:tbl[1]', namespaces=NS)[0] is node:
                            walk(nested)
                    return
                for child in node:
                    walk(child)

            walk(root)
    if revisions:
        issues.append('tracked_changes_present_current_view_not_certified_final')
    if not blocks:
        issues.append('no_extractable_text_possible_images_or_empty_document')
    return _result(blocks, 'needs_review' if revisions else ('parsed' if blocks else 'ocr_required'), issues,
                   revision_state='tracked_changes_present' if revisions else 'not_detected')


def _drawing_text(node):
    paragraphs = node.xpath('.//a:p', namespaces=NS)
    return '\n'.join(''.join(p.xpath('.//a:t/text()', namespaces=NS)) for p in paragraphs).strip()


def _bbox(node):
    transforms = node.xpath('./p:spPr/a:xfrm | ./p:xfrm | ./p:grpSpPr/a:xfrm', namespaces=NS)
    if not transforms:
        return None
    transform = transforms[0]
    off, ext = transform.find('a:off', NS), transform.find('a:ext', NS)
    if off is None or ext is None:
        return None
    return {'x': int(off.get('x', 0)), 'y': int(off.get('y', 0)),
            'width': int(ext.get('cx', 0)), 'height': int(ext.get('cy', 0)),
            'unit': 'EMU', 'coordinate_space': 'parent', 'rotation': transform.get('rot', '0')}


def parse_pptx(path, doc_id):
    blocks, issues = [], []
    with ZipFile(path) as z:
        slides = sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)),
                        key=lambda n: int(re.search(r'(\d+)\.xml$', n).group(1)))
        for part in slides:
            page = int(re.search(r'(\d+)\.xml$', part).group(1))
            root = _xml(z.read(part))
            for shape in root.xpath('//p:sp | //p:graphicFrame', namespaces=NS):
                identity = shape.xpath('./p:nvSpPr/p:cNvPr | ./p:nvGraphicFramePr/p:cNvPr', namespaces=NS)
                sid = identity[0].get('id') if identity else str(len(blocks))
                name = identity[0].get('name', '') if identity else ''
                base = {'doc_id': doc_id, 'location': {'part': part, 'slide': page, 'shape_id': sid, 'shape_name': name},
                        'bbox': _bbox(shape), 'section_path': []}
                groups = shape.xpath('ancestor::p:grpSp', namespaces=NS)
                if groups:
                    base['group_transforms'] = [_bbox(g) for g in groups]
                    issues.append('grouped_shapes_coordinates_in_parent_space')
                tables = shape.xpath('.//a:tbl', namespaces=NS)
                if tables:
                    for table_no, table in enumerate(tables, 1):
                        for row_no, row in enumerate(table.findall('a:tr', NS), 1):
                            cells = [_drawing_text(c) for c in row.findall('a:tc', NS)]
                            if not any(cells):
                                continue
                            blocks.append({**base, 'block_id': f'{doc_id}:{part}:s{sid}:t{table_no}r{row_no}',
                                           'kind': 'table_row', 'text': ' | '.join(cells), 'cells': cells,
                                           'location': {**base['location'], 'table': table_no, 'row': row_no}})
                else:
                    text = _drawing_text(shape)
                    if text:
                        blocks.append({**base, 'block_id': f'{doc_id}:{part}:s{sid}', 'kind': 'shape', 'text': text})
            if root.xpath('//a:blip', namespaces=NS):
                issues.append('embedded_images_not_ocr_processed')
            if root.xpath('//*[local-name()="chart"]', namespaces=NS):
                issues.append('embedded_charts_not_parsed_as_product_facts')
        if any(n.startswith('ppt/notesSlides/notesSlide') and n.endswith('.xml') for n in z.namelist()):
            issues.append('speaker_notes_excluded_from_display_body')
    if not blocks:
        issues.append('no_extractable_text_possible_images_or_empty_deck')
    return _result(blocks, 'parsed' if blocks else 'ocr_required', sorted(set(issues)))


def assess_pdf_text(text):
    compact = re.sub(r'\s+', '', text or '')
    if not compact:
        return {'status': 'no_text', 'char_count': 0, 'suspicious_count': 0}
    cid = re.findall(r'\(?cid:\d+\)?', compact, re.I)
    bad = sum(c == '\ufffd' or unicodedata.category(c) in {'Co', 'Cc'} for c in compact)
    visible = len(compact)
    damaged = sum(map(len, cid)) + bad
    return {'status': 'garbled_text' if damaged / visible > .02 or len(cid) > 1 else 'text_ok',
            'char_count': visible, 'suspicious_count': damaged}


def parse_pdf(path, doc_id):
    from pypdf import PdfReader
    blocks, issues = [], []
    with open(path, 'rb') as stream:
        pdf = PdfReader(stream, strict=False)
        if pdf.is_encrypted and not pdf.decrypt(''):
            return _result([], 'parse_error', ['encrypted_pdf_password_required'])
        for page_no, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ''
            quality = assess_pdf_text(text)
            blocks.append({'doc_id': doc_id, 'block_id': f'{doc_id}:page{page_no}', 'kind': 'pdf_page',
                           'text': text.strip(), 'location': {'page': page_no}, 'section_path': [],
                           'quality': quality['status'], 'quality_metrics': quality,
                           'bbox': {'x': 0, 'y': 0, 'width': float(page.mediabox.width), 'height': float(page.mediabox.height), 'unit': 'pt'}})
            if quality['status'] != 'text_ok':
                issues.append(f'page_{page_no}:{quality["status"]}')
    qualities = {b['quality'] for b in blocks}
    if qualities == {'no_text'} or not blocks:
        status = 'ocr_required'
    elif 'garbled_text' in qualities:
        status = 'garbled_text'
    elif 'no_text' in qualities:
        status = 'needs_review'
    else:
        status = 'parsed'
    return _result(blocks, status, issues)


def _cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None:
        return ''
    return str(value)


def parse_xlsx(path, doc_id, max_rows=30, max_columns=40):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    blocks, issues = [], []
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in wb:
            # Producer dimension metadata may be stale; reset rather than silently lose rows.
            original_dims = sheet.calculate_dimension(force=True)
            sheet.reset_dimensions()
            headers, count, found_more = [], 0, False
            for row in sheet.iter_rows(max_col=max_columns):
                if not any(c.value is not None for c in row):
                    continue
                if count >= max_rows:
                    found_more = True
                    break
                populated = [i for i, c in enumerate(row) if c.value is not None]
                end = max(populated) + 1
                row = row[:end]
                values = [_cell_value(c.value) for c in row]
                if not headers:
                    headers = values.copy()
                row_number = next(c.row for c in row if c.value is not None)
                formulas = {c.coordinate: c.value for c in row if c.data_type == 'f'}
                blocks.append({'doc_id': doc_id, 'block_id': f'{doc_id}:sheet:{sheet.title}:r{row_number}',
                               'kind': 'spreadsheet_row', 'text': ' | '.join(values), 'cells': values,
                               'headers': headers.copy(), 'header_status': 'first_nonempty_row_candidate',
                               'formulas': formulas, 'section_path': [sheet.title],
                               'location': {'sheet': sheet.title, 'row': row_number,
                                            'cell_range': f'A{row_number}:{get_column_letter(end)}{row_number}'}})
                count += 1
            issues.append(f'sheet:{sheet.title}:sample_limit={max_rows}_rows,{max_columns}_columns;dimensions={original_dims};more_rows={found_more}')
        if any(b.get('formulas') for b in blocks):
            issues.append('formulas_preserved_not_recalculated_or_treated_as_numeric_values')
    finally:
        wb.close()
    return _result(blocks, 'sampled', issues)


def parse_document(path, doc_id, xlsx_max_rows=30):
    path = Path(path)
    try:
        if path.suffix.lower() == '.docx':
            return parse_docx(path, doc_id)
        if path.suffix.lower() == '.pptx':
            return parse_pptx(path, doc_id)
        if path.suffix.lower() == '.pdf':
            return parse_pdf(path, doc_id)
        if path.suffix.lower() == '.xlsx':
            return parse_xlsx(path, doc_id, max_rows=xlsx_max_rows)
        if path.suffix.lower() in {'.png', '.jpg', '.jpeg'}:
            from PIL import Image
            with Image.open(path) as img:
                return _result([], 'ocr_required', [f'image_requires_ocr:{img.width}x{img.height}'])
        if path.suffix.lower() in {'.doc', '.xls'}:
            return _result([], 'legacy_pending', ['legacy_binary_requires_read_only_conversion'])
        return _result([], 'unsupported', ['unsupported_extension'])
    except Exception as exc:
        return _result([], 'parse_error', [f'{type(exc).__name__}: {str(exc)[:300]}'])


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _cache_items(project_dir, name):
    index = Path(project_dir) / 'cache' / name / 'index.json'
    if not index.exists():
        return []
    value = json.loads(index.read_text(encoding='utf-8-sig'))
    if isinstance(value, dict): value = value.get('items', [])
    if not isinstance(value, list): raise ValueError('cache index must contain a list of items')
    return value


def _cache_path(value, project_dir):
    """Resolve portable cache references independently of the caller's cwd."""
    path = Path(value)
    return (path if path.is_absolute() else Path(project_dir) / path).resolve()


def merge_ocr_blocks(native, cached_blocks, doc_id):
    """Replace only bad/native-missing pages; never overwrite healthy native text."""
    good_pages = {b['location'].get('page') for b in native['blocks'] if b.get('quality') == 'text_ok'}
    accepted = [b for b in cached_blocks if b.get('location', {}).get('page') not in good_pages
                and b.get('text', '').strip() and b.get('location', {}).get('page') is not None]
    replaced = {b['location']['page'] for b in accepted}
    blocks = [b for b in native['blocks'] if b['location'].get('page') not in replaced]
    for index, block in enumerate(accepted, 1):
        blocks.append({**block, 'doc_id': doc_id,
                       'block_id': f'{doc_id}:ocr:p{block["location"]["page"]}:{index}',
                       'section_path': block.get('section_path', []), 'quality': 'ocr_unverified',
                       'evidence_method': 'ocr_cache_not_human_verified'})
    blocks.sort(key=lambda b: (b['location'].get('page', 0), b['location'].get('line', 0)))
    return {**native, 'blocks': blocks, 'parse_status': 'ocr_unverified' if accepted else native['parse_status'],
            'issues': native['issues'] + (['ocr_cache_used_unverified;original_quality_issues_retained'] if accepted else ['ocr_cache_no_usable_bad_page_text'])}


def parse_document_cached(path, doc_id, source_sha256, project_dir, xlsx_max_rows=30):
    path = Path(path)
    kind = 'legacy' if path.suffix.lower() in {'.doc', '.xls'} else ('ocr' if path.suffix.lower() in {'.pdf', '.png', '.jpg', '.jpeg'} else None)
    cache = None
    cache_issues = []
    if kind:
        try:
            for entry in _cache_items(project_dir, kind):
                if str(_cache_path(entry['source'], project_dir)).casefold() == str(path.resolve()).casefold():
                    if entry.get('sha256') != source_sha256:
                        cache_issues.append(f'{kind}_cache_hash_mismatch_ignored')
                    else:
                        cache = entry
                    break
        except (OSError, ValueError, KeyError, TypeError) as exc:
            cache_issues.append(f'{kind}_cache_unreadable:{type(exc).__name__}')
    if cache and kind == 'legacy':
        converted_digest = cache.get('converted_sha256')
        if not converted_digest:
            cache_issues.append('legacy_cache_converted_hash_missing_ignored')
        else:
            try:
                converted = _cache_path(cache['converted'], project_dir)
                if _sha256(path) != source_sha256:
                    cache_issues.append('legacy_cache_source_hash_mismatch_ignored')
                elif _sha256(converted) != converted_digest:
                    cache_issues.append('legacy_cache_converted_hash_mismatch_ignored')
                else:
                    result = parse_document(converted, doc_id, xlsx_max_rows=xlsx_max_rows)
                    # Discard evidence if either file changed while it was read.
                    if _sha256(path) != source_sha256:
                        cache_issues.append('legacy_cache_source_hash_mismatch_ignored')
                    elif _sha256(converted) != converted_digest:
                        cache_issues.append('legacy_cache_converted_hash_mismatch_ignored')
                    else:
                        for block in result['blocks']:
                            block['location']['converted_copy'] = str(converted)
                            block['evidence_method'] = 'converted_copy'
                        if result['parse_status'] == 'parsed':
                            result['parse_status'] = 'converted_copy_parsed'
                        result['issues'].append('legacy_source_parsed_from_sha256_matched_read_only_copy')
                        return result
            except (OSError, ValueError, KeyError, TypeError) as exc:
                cache_issues.append(f'legacy_cache_conversion_unreadable:{type(exc).__name__}')
    result = parse_document(path, doc_id, xlsx_max_rows=xlsx_max_rows)
    result['issues'].extend(cache_issues)
    if cache and kind == 'ocr':
        blocks = cache.get('blocks', [])
        if cache.get('blocks_path'):
            blocks = json.loads(Path(cache['blocks_path']).read_text(encoding='utf-8-sig'))
        result = merge_ocr_blocks(result, blocks, doc_id)
    return result


def enrich_corpus(project_dir, refresh_extensions=()):
    """Apply validated new legacy/OCR caches without rescanning unaffected sources."""
    project_dir = Path(project_dir)
    data_dir = project_dir / 'data'
    manifest_path = data_dir / 'manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    cached_sources = {str(_cache_path(item['source'], project_dir)).casefold()
                      for kind in ('legacy', 'ocr') for item in _cache_items(project_dir, kind)}
    replacements, changed = {}, []
    for entry in manifest:
        if entry['parse_status'] == 'excluded' or (str(Path(entry['absolute_path']).resolve()).casefold() not in cached_sources
                and entry['extension'] not in refresh_extensions):
            continue
        # Re-hash live input before attaching cached evidence to prevent a stale parse join.
        if _sha256(entry['absolute_path']) != entry['sha256']:
            entry['issues'].append('live_source_hash_changed_cache_not_attached')
            entry['parse_status'] = 'needs_review'
            continue
        parsed = parse_document_cached(entry['absolute_path'], entry['doc_id'], entry['sha256'], project_dir)
        blocks = parsed.pop('blocks')
        entry.update(parsed)
        entry.update(classify_document(entry['relative_path'], blocks))
        entry['block_count'] = len(blocks)
        replacements[entry['doc_id']] = blocks
        changed.append(entry['doc_id'])
    block_path = data_dir / 'blocks.jsonl'
    temporary = data_dir / 'corpus_blocks.enrich.tmp'
    with block_path.open(encoding='utf-8') as src, temporary.open('w', encoding='utf-8') as dst:
        for line in src:
            if json.loads(line)['doc_id'] not in replacements:
                dst.write(line)
        for blocks in replacements.values():
            for block in blocks:
                dst.write(json.dumps(block, ensure_ascii=False) + '\n')
    temporary.replace(block_path)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    summary_path = data_dir / 'corpus_summary.json'
    summary = json.loads(summary_path.read_text(encoding='utf-8'))
    active = [d for d in manifest if d['parse_status'] != 'excluded']
    summary.update({'status_counts': dict(Counter(d['parse_status'] for d in manifest)),
                    'material_counts': dict(Counter(d['material_kind'] for d in active)),
                    'role_counts': dict(Counter(d['document_role'] for d in active)),
                    'block_count': sum(d['block_count'] for d in manifest),
                    'cache_or_parser_refreshed_documents': len(changed)})
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'manifest': manifest, 'summary': summary}


def exclusion_reason(relative, project_relative):
    path = Path(relative)
    pieces = path.parts
    if project_relative and (path == project_relative or project_relative in path.parents):
        return 'project_output'
    for piece in pieces[:-1]:
        low = piece.lower()
        if low.startswith('.'):
            return 'hidden_directory'
        if low in EXCLUDED_DIRS or low.startswith(('_work', '_tmp', '_proof', '_qa', 'corpus_')) or low.endswith(('_build', '_qa')):
            return 'build_qa_or_output_directory'
    if path.name.startswith(('~$', '.')):
        return 'temporary_or_hidden_file'
    return None


def classify_document(relative_path, blocks):
    # OCR often inserts spaces between every Chinese character. Normalize only
    # the search view; preserve each source block and its evidence text unchanged.
    text = re.sub(r'[^\S\r\n]+', '', '\n'.join(b['text'] for b in blocks[:100]))
    early = re.sub(r'[^\S\r\n]+', '', '\n'.join(b['text'] for b in blocks[:20]))
    path = str(relative_path).replace('\\', '/')
    name = Path(path).name
    nonproduct = bool(re.search(r'(?:^|/)(?:排名|01_FICC周报)(?:/|$)|考勤|出勤|工作小结|工作总结|采购|证明材料|答辩|身份证|通讯录|校友|简历', path, re.I))
    kind = 'unknown'
    # Prefer explicit material names, then document-title/body evidence; never certify identity.
    patterns = [('contract', r'合同|信托文件'), ('prospectus', r'招募说明书|产品资料概要'),
                ('channel_calendar', r'产品日历|日历录入|渠道.{0,8}模板'), ('onepager', r'一页通'),
                ('poster', r'海报|四宫格|长图'), ('qa', r'Q[＆&]?A|问答'),
                ('roadshow', r'路演|营销垫板'), ('product_highlights', r'产品亮点')]
    for candidate, pattern in patterns:
        if re.search(pattern, name, re.I):
            kind = candidate
            break
    if kind == 'unknown':
        if re.search(r'(?:管理|资产管理|信托|基金)合同', early):
            kind = 'contract'
        elif re.search(r'招募说明书', early):
            kind = 'prospectus'
        elif Path(name).suffix.lower() == '.pptx' and re.search(r'产品要素|投资策略|投资范围', text):
            kind = 'presentation'
        elif Path(name).suffix.lower() in {'.png','.jpg','.jpeg'} and re.search(r'产品要素|产品名称|基金名称|认购|申购费|投资策略|风险等级', text):
            kind = 'image_material'
    channels = [('建设银行', r'建行|建设银行'), ('招商银行', r'招行|招商银行'), ('工商银行', r'工行|工商银行'),
                ('中国银行', r'中行|中国银行'), ('民生银行', r'民生'), ('上海银行', r'上海银行'),
                ('东莞银行', r'东莞银行'), ('邮储银行', r'邮储'), ('苏州银行', r'苏州银行'),
                ('平安银行', r'平安(?:银行|版|海报)'), ('集团', r'集团版'), ('通用', r'通用版')]
    channel, channel_source = 'unknown', 'not_identified'
    for candidate, pattern in channels:
        if re.search(pattern, name):
            channel, channel_source = candidate, 'filename_candidate'
            break
    if channel == 'unknown':
        for candidate, pattern in channels:
            if re.search(r'(?:渠道|销售机构|代销机构)[：:\s]*(?:' + pattern + ')', early) or re.search(r'(?:' + pattern + r')(?:作为|为)(?:本[^，。\n]{0,16})?(?:代销|销售)机构', early):
                channel, channel_source = candidate, 'body_candidate'
                break
    identities = []
    seen = set()
    product_pattern = re.compile(r'(?:国泰(?:海通|君安)|海通|平安信托|中信信托|中粮信托|外贸信托)[^\n|。；：:]{2,85}?(?:集合资产管理计划|单一资产管理计划|资产管理信托|集合资金信托计划|证券投资基金|FOF[0-9Xx]*号|QDII[0-9Xx]+号)')
    for block in blocks[:100]:
        for hit in product_pattern.finditer(re.sub(r'[^\S\r\n]+', '', block['text'])):
            value = hit.group(0).strip()
            if value not in seen:
                identities.append({'value': value, 'source': 'body', 'confidence': 'candidate',
                                   'block_id': block['block_id'], 'location': block['location']})
                seen.add(value)
        if len(identities) >= 5:
            break
    if not identities and not nonproduct:
        identities.append({'value': Path(name).stem, 'source': 'filename', 'confidence': 'low_unverified'})
    role = 'non_product' if nonproduct else ('product_candidate' if kind != 'unknown' or any(c['source'] == 'body' for c in identities) else 'unknown')
    return {'material_kind': kind, 'material_kind_status': 'candidate' if kind != 'unknown' else 'unknown',
            'channel': channel, 'channel_source': channel_source, 'document_role': role,
            'identity_candidates': identities[:5], 'identity_confirmed': False}


def scan_corpus(root, project_dir, write_outputs=True, xlsx_max_rows=30):
    root, project_dir = Path(root).resolve(), Path(project_dir).resolve()
    try:
        project_relative = project_dir.relative_to(root)
    except ValueError:
        project_relative = None
    manifest, all_blocks, other_files, excluded_other_files, walk_errors = [], [], Counter(), Counter(), []
    excluded_subtrees = []
    data_dir = project_dir / 'data'
    if write_outputs:
        data_dir.mkdir(parents=True, exist_ok=True)
    stream = (data_dir / 'blocks.jsonl').open('w', encoding='utf-8') if write_outputs else None
    try:
        for directory, dirs, files in os.walk(root, followlinks=False, onerror=lambda err: walk_errors.append(str(err))):
            dirs.sort(key=str.casefold)
            kept_dirs = []
            for child_name in dirs:
                child = Path(directory) / child_name
                relative_child = child.relative_to(root)
                reason = None
                if relative_child == project_relative:
                    reason = 'project_output'
                elif child_name.casefold() == 'node_modules':
                    reason = 'dependency_directory'
                else:
                    try:
                        if child.is_symlink() or getattr(child.lstat(), 'st_file_attributes', 0) & 0x400:
                            reason = 'reparse_or_symlink_directory'
                    except OSError:
                        reason = 'unreadable_directory'
                if reason:
                    excluded_subtrees.append({'relative_path': relative_child.as_posix(), 'reason': reason,
                                              'contents_enumerated': False})
                else:
                    kept_dirs.append(child_name)
            dirs[:] = kept_dirs
            for filename in sorted(files, key=str.casefold):
                path = Path(directory) / filename
                relative = path.relative_to(root)
                reason = exclusion_reason(relative, project_relative)
                if path.suffix.lower() not in SUPPORTED:
                    counter = excluded_other_files if reason else other_files
                    counter[path.suffix.lower() or '[no extension]'] += 1
                    continue
                doc_id = hashlib.sha256(relative.as_posix().encode('utf-8')).hexdigest()[:20]
                entry = {'doc_id': doc_id, 'relative_path': relative.as_posix(), 'absolute_path': str(path),
                         'sha256': None, 'extension': path.suffix.lower(), 'block_count': 0, 'issues': []}
                try:
                    before = path.stat()
                    entry['sha256'] = _sha256(path)
                    entry['size_bytes'] = before.st_size
                    entry['mtime_ns'] = before.st_mtime_ns
                except OSError as exc:
                    result = _result([], 'parse_error', [f'source_read_error:{type(exc).__name__}'])
                else:
                    if reason:
                        result = _result([], 'excluded', [reason])
                    else:
                        result = parse_document_cached(path, doc_id, entry['sha256'], project_dir, xlsx_max_rows=xlsx_max_rows)
                        try:
                            after = path.stat()
                            if before.st_size != after.st_size or before.st_mtime_ns != after.st_mtime_ns:
                                result['issues'].append('source_changed_during_scan')
                                result['parse_status'] = 'needs_review'
                        except OSError:
                            result['issues'].append('source_unavailable_after_parse')
                blocks = result.pop('blocks')
                entry.update(result)
                entry['block_count'] = len(blocks)
                entry.update(classify_document(relative.as_posix(), blocks))
                if reason:
                    entry['document_role'] = 'excluded'
                manifest.append(entry)
                if stream:
                    for block in blocks:
                        stream.write(json.dumps(block, ensure_ascii=False) + '\n')
                else:
                    all_blocks.extend(blocks)
                if write_outputs and len(manifest) % 100 == 0:
                    print(json.dumps({'inventoried': len(manifest), 'blocks': sum(d['block_count'] for d in manifest)}, ensure_ascii=False), flush=True)
    finally:
        if stream:
            stream.close()
    active = [d for d in manifest if d['parse_status'] != 'excluded']
    summary = {'schema_version': 1, 'root': str(root), 'generated_at': datetime.now(timezone.utc).isoformat(),
               'supported_files': len(manifest), 'included_files': len(active),
               'excluded_files': len(manifest) - len(active), 'block_count': sum(d['block_count'] for d in manifest),
               'status_counts': dict(Counter(d['parse_status'] for d in manifest)),
               'included_extension_counts': dict(Counter(d['extension'] for d in active)),
               'material_counts': dict(Counter(d['material_kind'] for d in active)),
               'role_counts': dict(Counter(d['document_role'] for d in active)),
               'exclusion_counts': dict(Counter(i for d in manifest if d['parse_status'] == 'excluded' for i in d['issues'])),
               'other_extension_counts': dict(other_files), 'excluded_other_extension_counts': dict(excluded_other_files),
               'excluded_subtrees': excluded_subtrees, 'walk_errors': walk_errors,
               'boundaries': ['Excluded supported documents are itemized but not parsed.',
                              f'XLSX: first {xlsx_max_rows} nonempty rows and 40 columns per sheet; headers are candidates.',
                              'Dependency, reparse/junction and this project subtrees are logged and pruned before traversal.',
                              'DOC/XLS use read-only converted copies; product images require OCR; archives counted only.',
                              'All identities and channel classifications are candidates, not certified.',
                              'Tracked changes use a proposed current view and always require review.']}
    if write_outputs:
        (data_dir / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        (data_dir / 'corpus_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'manifest': manifest, 'blocks': all_blocks, 'summary': summary}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument('--project-dir', type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument('--xlsx-max-rows', type=int, default=30)
    parser.add_argument('--enrich-only', action='store_true', help='Attach validated legacy/OCR caches to existing inventory')
    args = parser.parse_args()
    logging.getLogger('pypdf').setLevel(logging.ERROR)
    result = enrich_corpus(args.project_dir) if args.enrich_only else scan_corpus(args.root, args.project_dir, xlsx_max_rows=args.xlsx_max_rows)
    print(json.dumps(result['summary'], ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
